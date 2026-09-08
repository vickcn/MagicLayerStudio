from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union


DEFAULT_SLIDE_WIDTH_IN = 13.333333
EMU_PER_INCH = 914400
PT_PER_INCH = 72.0

# 重建模式
MODE_IMAGE_LAYER = "image_layer"  # 透明 PNG 貼圖
MODE_WORDART    = "wordart"       # PPTX TextBox 物件
MODE_MIXED      = "mixed"         # 僅 wordart；無法推算樣式時 fallback 至 image_layer


def rebuild_pptx_from_document(
    document_json_path: Path,
    output_pptx_path: Path,
    rebuild_mode: str = MODE_IMAGE_LAYER,
    custom_edits: Optional[Union[dict, Path, str]] = None,
) -> Path:
    try:
        import pptx  # noqa: F401
    except ImportError:
        return rebuild_pptx_with_padocr(
            document_json_path,
            output_pptx_path,
            rebuild_mode=rebuild_mode,
            custom_edits=custom_edits,
        )

    return _rebuild_pptx(
        document_json_path,
        output_pptx_path,
        rebuild_mode=rebuild_mode,
        custom_edits=custom_edits,
    )


def rebuild_pptx_with_padocr(
    document_json_path: Path,
    output_pptx_path: Path,
    rebuild_mode: str = MODE_IMAGE_LAYER,
    custom_edits: Optional[Union[dict, Path, str]] = None,
) -> Path:
    conda_exe = os.environ.get("CONDA_EXE") or shutil.which("conda")
    if conda_exe is None:
        raise RuntimeError(
            "python-pptx is not available. Activate padocr or make conda available."
        )

    cmd = [
        conda_exe,
        "run",
        "--no-capture-output",
        "-n",
        "padocr",
        "python",
        "-m",
        "src.pptx_rebuilder",
        str(document_json_path),
        str(output_pptx_path),
        "--mode",
        rebuild_mode,
    ]
    if custom_edits:
        if isinstance(custom_edits, (str, Path)):
            cmd.extend(["--custom-edits", str(custom_edits)])
        elif isinstance(custom_edits, dict):
            temp_edits_file = output_pptx_path.parent / "_temp_custom_edits.json"
            with open(temp_edits_file, "w", encoding="utf-8") as f:
                json.dump(custom_edits, f, ensure_ascii=False)
            cmd.extend(["--custom-edits", str(temp_edits_file)])

    subprocess.run(cmd, check=True)
    return output_pptx_path


def _resolve_file_path(base_dir: Path, path_val: str | Path) -> Path:
    p = Path(str(path_val))
    if p.is_absolute() and p.exists():
        return p
    if p.exists():
        return p.resolve()
    if (base_dir / p).exists():
        return (base_dir / p).resolve()
    if (base_dir / p.name).exists():
        return (base_dir / p.name).resolve()
    return base_dir / p


def _rebuild_pptx(
    document_json_path: Path,
    output_pptx_path: Path,
    rebuild_mode: str = MODE_IMAGE_LAYER,
    custom_edits: Optional[Union[dict, Path, str]] = None,
) -> Path:
    from pptx import Presentation
    from pptx.util import Emu

    document_json_path = document_json_path.resolve()
    output_pptx_path = output_pptx_path.resolve()
    output_dir = document_json_path.parent

    with open(document_json_path, "r", encoding="utf-8") as file:
        document = json.load(file)

    pages = document.get("pages", [])
    if not pages:
        raise RuntimeError(f"No pages found in: {document_json_path}")

    # 載入自訂 edits 資料（若傳入路徑或字典）
    edits_map: dict = {}
    if isinstance(custom_edits, (str, Path)):
        edits_path = Path(custom_edits)
        if edits_path.exists():
            try:
                with open(edits_path, "r", encoding="utf-8") as f:
                    edits_map = json.load(f)
            except Exception:
                edits_map = {}
    elif isinstance(custom_edits, dict):
        edits_map = custom_edits

    first_objects = _load_objects_json(output_dir, pages[0])
    canvas_width = int(first_objects["canvas"]["width"])
    canvas_height = int(first_objects["canvas"]["height"])
    slide_width = int(round(DEFAULT_SLIDE_WIDTH_IN * EMU_PER_INCH))
    slide_height = int(round(slide_width * canvas_height / canvas_width))

    prs = Presentation()
    prs.slide_width = Emu(slide_width)
    prs.slide_height = Emu(slide_height)
    blank_layout = prs.slide_layouts[6]

    for p_idx, page in enumerate(pages):
        objects_json = _load_objects_json(output_dir, page)
        page_dir_val = page.get("output_dir") or page.get("output", "")
        page_dir = _resolve_file_path(output_dir, page_dir_val)
        page_canvas_width = int(objects_json["canvas"]["width"])
        page_canvas_height = int(objects_json["canvas"]["height"])

        # 取得此頁的 custom edits（支援 key 為 "0" 或 0 或 page.id）
        page_id = page.get("page_id") or page.get("id") or str(p_idx)
        page_edits = edits_map.get(str(p_idx)) or edits_map.get(p_idx) or edits_map.get(page_id) or {}

        slide = prs.slides.add_slide(blank_layout)
        background_path = _resolve_background_path(output_dir, page)
        slide.shapes.add_picture(
            str(background_path),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )

        # 讀取 layers.json（供 wordart 模式使用）
        layers_by_file = _load_layers_index(output_dir, page)
        handled_edit_keys = set()

        for obj in objects_json.get("objects", []):
            image_path = _resolve_file_path(page_dir, obj.get("file", ""))
            if not image_path.exists():
                if "image_path" in obj:
                    image_path = _resolve_file_path(page_dir, obj["image_path"])
                if not image_path.exists():
                    continue

            # 檢查此物件或其 layer_ids 是否有自訂 edit
            obj_id = obj.get("id", "")
            layer_ids = obj.get("layer_ids", [])
            custom_item = page_edits.get(obj_id)
            if custom_item:
                handled_edit_keys.add(obj_id)
            else:
                for lid in layer_ids:
                    if lid in page_edits:
                        custom_item = page_edits[lid]
                        handled_edit_keys.add(lid)
                        break

            # 若使用者刪除了該物件，則略過不輸出
            if custom_item and custom_item.get("deleted"):
                continue

            # 決定座標（若有自訂 x, y, width, height 則覆寫）
            pos_x = custom_item["x"] if (custom_item and "x" in custom_item) else int(obj["x"])
            pos_y = custom_item["y"] if (custom_item and "y" in custom_item) else int(obj["y"])
            pos_w = custom_item["width"] if (custom_item and "width" in custom_item) else int(obj["width"])
            pos_h = custom_item["height"] if (custom_item and "height" in custom_item) else int(obj["height"])

            left = _scale_emu(int(pos_x), page_canvas_width, prs.slide_width)
            top = _scale_emu(int(pos_y), page_canvas_height, prs.slide_height)
            width = _scale_emu(int(pos_w), page_canvas_width, prs.slide_width)
            height = _scale_emu(int(pos_h), page_canvas_height, prs.slide_height)

            # 決定模式（若自訂有指定 mode，以自訂為準）
            if custom_item and "mode" in custom_item:
                use_wordart = (custom_item["mode"] == MODE_WORDART)
            else:
                use_wordart = _should_use_wordart(
                    rebuild_mode,
                    obj,
                    image_path,
                    layers_by_file,
                )

            if use_wordart:
                layer_info = layers_by_file.get(str(image_path), {})
                _add_wordart_textbox(
                    slide, prs, obj, layer_info,
                    left, top, width, height,
                    page_canvas_width, page_canvas_height,
                    custom_item=custom_item,
                )
            else:
                slide.shapes.add_picture(
                    str(image_path),
                    left,
                    top,
                    width=width,
                    height=height,
                )

        # 處理使用者新增或複製出來的全新文字物件
        for edit_key, new_item in page_edits.items():
            if edit_key in handled_edit_keys:
                continue
            if not isinstance(new_item, dict) or new_item.get("deleted"):
                continue

            pos_x = int(new_item.get("x", 100))
            pos_y = int(new_item.get("y", 100))
            pos_w = int(new_item.get("width", 400))
            pos_h = int(new_item.get("height", 100))

            left = _scale_emu(pos_x, page_canvas_width, prs.slide_width)
            top = _scale_emu(pos_y, page_canvas_height, prs.slide_height)
            width = _scale_emu(pos_w, page_canvas_width, prs.slide_width)
            height = _scale_emu(pos_h, page_canvas_height, prs.slide_height)

            _add_wordart_textbox(
                slide, prs, new_item, {},
                left, top, width, height,
                page_canvas_width, page_canvas_height,
                custom_item=new_item,
            )

    output_pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_pptx_path))
    return output_pptx_path


def _load_objects_json(output_dir: Path, page: dict) -> dict:
    objects_path_val = page.get("objects_json", "")
    objects_path = _resolve_file_path(output_dir, objects_path_val)
    with open(objects_path, "r", encoding="utf-8") as file:
        return json.load(file)


def _load_layers_index(output_dir: Path, page: dict) -> Dict[str, dict]:
    """
    讀取 layers.json，回傳 {abs_file_path_str: layer_info} 的字典。
    """
    layers_path_val = page.get("layers_json", "")
    if not layers_path_val:
        return {}
    layers_path = _resolve_file_path(output_dir, layers_path_val)
    if not layers_path.exists():
        return {}
    try:
        with open(layers_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        index = {}
        for layer in data.get("layers", []):
            file_val = layer.get("file", "")
            abs_path = _resolve_file_path(layers_path.parent, file_val)
            index[str(abs_path)] = layer
        return index
    except Exception:
        return {}


def _should_use_wordart(
    rebuild_mode: str,
    obj: dict,
    image_path: Path,
    layers_by_file: Dict[str, dict],
) -> bool:
    if rebuild_mode == MODE_IMAGE_LAYER:
        return False
    layer_info = layers_by_file.get(str(image_path), {})
    has_style_hint = bool(layer_info.get("style_hint"))
    has_text = bool(layer_info.get("text", "").strip())
    if rebuild_mode == MODE_WORDART:
        return has_style_hint and has_text
    if rebuild_mode == MODE_MIXED:
        return has_style_hint and has_text
    return False


def _add_wordart_textbox(
    slide,
    prs,
    obj: dict,
    layer_info: dict,
    left,
    top,
    width,
    height,
    canvas_width: int,
    canvas_height: int,
    custom_item: Optional[dict] = None,
) -> None:
    """以 python-pptx TextBox 模擬 WordArt 樣式展示文字（支援自訂樣式覆寫）。"""
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    from pptx.enum.text import PP_ALIGN

    # 文字優先順序：custom_item["text"] > obj["text"] > layer_info["text"]
    text = ""
    if custom_item and "text" in custom_item:
        text = str(custom_item["text"]).strip()
    if not text:
        text = obj.get("text", "").strip() or layer_info.get("text", "").strip()
    if not text:
        return

    style = layer_info.get("style_hint", {})
    custom_style = (custom_item.get("style") or {}) if custom_item else {}

    # 樣式覆寫
    color_rgb = custom_style.get("color_rgb") or style.get("dominant_color_rgb", [255, 255, 255])
    if isinstance(color_rgb, list) and len(color_rgb) == 3:
        pass
    else:
        color_rgb = [255, 255, 255]

    font_size_pt = float(custom_style.get("font_size_pt") or style.get("estimated_font_size_pt", 24.0))
    likely_bold = bool(custom_style.get("bold") if "bold" in custom_style else style.get("likely_bold", False))
    font_italic = bool(custom_style.get("italic", False))
    font_name = custom_style.get("font_name") or style.get("font_name_hint") or "Noto Sans TC"

    # 對齊方式
    align_str = custom_style.get("align", "center").lower()
    align_map = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }
    pp_align = align_map.get(align_str, PP_ALIGN.CENTER)

    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    # 處理多行文字
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = pp_align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size_pt)
        run.font.bold = likely_bold
        run.font.italic = font_italic
        run.font.name = font_name
        try:
            run.font.color.rgb = RGBColor(int(color_rgb[0]), int(color_rgb[1]), int(color_rgb[2]))
        except Exception:
            pass


def _resolve_background_path(output_dir: Path, page: dict) -> Path:
    background = page.get("background") or page.get("background_path")
    if background:
        background_path = _resolve_file_path(output_dir, background)
        if background_path.exists():
            return background_path

    page_obj = page.get("page")
    source_image = (
        page.get("source_image")
        or (page_obj.get("image_path") if isinstance(page_obj, dict) else None)
    )
    if not source_image:
        raise RuntimeError(f"Page has no background or source image: {page}")

    source_path = _resolve_file_path(output_dir, source_image)
    if source_path.exists():
        return source_path

    raise RuntimeError(f"Cannot resolve page background: {source_image}")


def _scale_emu(value: int, source_size: int, target_size: int):
    from pptx.util import Emu

    return Emu(int(round(value / source_size * int(target_size))))


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("document_json", type=Path)
    parser.add_argument("output_pptx", type=Path)
    parser.add_argument(
        "--mode",
        choices=[MODE_IMAGE_LAYER, MODE_WORDART, MODE_MIXED],
        default=MODE_IMAGE_LAYER,
        help="重建模式：image_layer（貼圖）/ wordart（文字物件）/ mixed",
    )
    parser.add_argument(
        "--custom-edits",
        type=Path,
        default=None,
        help="自訂圖層編輯 JSON 檔路徑",
    )
    args = parser.parse_args(argv)
    rebuilt_path = rebuild_pptx_from_document(
        args.document_json,
        args.output_pptx,
        rebuild_mode=args.mode,
        custom_edits=args.custom_edits,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
