"""
pptx_font_reader.py
從原始 PPTX 讀取每頁文字框的字型資訊，並以座標重疊度
對映到 OCR 結果的 bounding box，為 style_hint 補上
font_name_hint。

需求：python-pptx（padocr conda 環境）
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional, Tuple

try:
    from pptx import Presentation
    from pptx.util import Emu
    _PPTX_AVAILABLE = True
except ImportError:
    _PPTX_AVAILABLE = False


# ── 公開 API ──────────────────────────────────────────────────────────────────

def enrich_layers_with_pptx_fonts(
    layers: List[dict],
    pptx_path: Path,
    slide_index: int,
    canvas_width: int,
    canvas_height: int,
) -> None:
    """
    就地修改 layers 列表中每個 layer 的 style_hint.font_name_hint，
    從 PPTX 原始文字框依座標對映推算字體名稱。

    Parameters
    ----------
    layers : List[dict]
        text_layer_extractor 回傳的 layers 列表（就地修改）
    pptx_path : Path
        原始 PPTX 路徑
    slide_index : int
        0-based 投影片索引
    canvas_width / canvas_height : int
        渲染圖像的像素尺寸（用於座標換算）
    """
    if not _PPTX_AVAILABLE:
        return
    if not layers:
        return

    try:
        font_shapes = _extract_font_shapes(pptx_path, slide_index)
    except Exception:
        return

    if not font_shapes:
        return

    try:
        prs = Presentation(str(pptx_path))
        slide_w = int(prs.slide_width)
        slide_h = int(prs.slide_height)
    except Exception:
        return

    for layer in layers:
        ocr_box = layer.get("ocr_box")
        if not ocr_box:
            continue

        # OCR box 轉換成 EMU 座標
        x1_px, y1_px, x2_px, y2_px = ocr_box
        x1_emu = int(x1_px / canvas_width * slide_w)
        y1_emu = int(y1_px / canvas_height * slide_h)
        x2_emu = int(x2_px / canvas_width * slide_w)
        y2_emu = int(y2_px / canvas_height * slide_h)

        best_font = _find_best_font(
            x1_emu, y1_emu, x2_emu, y2_emu,
            font_shapes,
        )
        if best_font:
            if "style_hint" not in layer:
                layer["style_hint"] = {}
            layer["style_hint"]["font_name_hint"] = best_font


# ── 內部函式 ──────────────────────────────────────────────────────────────────

def _extract_font_shapes(
    pptx_path: Path,
    slide_index: int,
) -> List[Dict]:
    """
    從 PPTX 讀取指定投影片上所有文字框的幾何與字型資訊。

    回傳格式：
    [
        {
            "left": int,   # EMU
            "top": int,
            "right": int,
            "bottom": int,
            "font_name": str | None,
            "font_size_pt": float | None,
        }, ...
    ]
    """
    prs = Presentation(str(pptx_path))
    slides = prs.slides
    if slide_index >= len(slides):
        return []

    slide = slides[slide_index]
    shapes_info = []

    for shape in slide.shapes:
        left = int(getattr(shape, "left", 0) or 0)
        top = int(getattr(shape, "top", 0) or 0)
        w = int(getattr(shape, "width", 0) or 0)
        h = int(getattr(shape, "height", 0) or 0)
        right = left + w
        bottom = top + h

        if w <= 0 or h <= 0:
            continue

        font_name = None
        font_size_pt = None

        # 文字框
        if getattr(shape, "has_text_frame", False):
            tf = shape.text_frame
            for para in tf.paragraphs:
                for run in para.runs:
                    f = run.font
                    if f.name:
                        font_name = f.name
                    if f.size is not None:
                        font_size_pt = float(f.size.pt)
                    if font_name:
                        break
                if font_name:
                    break

        # 表格
        elif getattr(shape, "has_table", False):
            try:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                if run.font.name:
                                    font_name = run.font.name
                                    break
                            if font_name:
                                break
                        if font_name:
                            break
                    if font_name:
                        break
            except Exception:
                pass

        shapes_info.append({
            "left": left,
            "top": top,
            "right": right,
            "bottom": bottom,
            "font_name": font_name,
            "font_size_pt": font_size_pt,
        })

    return shapes_info


def _iou(
    ax1: int, ay1: int, ax2: int, ay2: int,
    bx1: int, by1: int, bx2: int, by2: int,
) -> float:
    """計算兩個矩形的 IoU（Intersection over Union）。"""
    ix1 = max(ax1, bx1)
    iy1 = max(ay1, by1)
    ix2 = min(ax2, bx2)
    iy2 = min(ay2, by2)

    if ix2 <= ix1 or iy2 <= iy1:
        return 0.0

    inter = (ix2 - ix1) * (iy2 - iy1)
    area_a = max(0, (ax2 - ax1) * (ay2 - ay1))
    area_b = max(0, (bx2 - bx1) * (by2 - by1))
    union = area_a + area_b - inter
    return inter / union if union > 0 else 0.0


def _find_best_font(
    x1: int, y1: int, x2: int, y2: int,
    font_shapes: List[Dict],
    min_iou: float = 0.10,
) -> Optional[str]:
    """
    找出與 OCR box 重疊度最高的 PPTX shape，回傳其字型名稱。
    IoU 低於 min_iou 則視為未匹配。
    """
    best_iou = 0.0
    best_font: Optional[str] = None

    for shape in font_shapes:
        if not shape.get("font_name"):
            continue
        iou = _iou(
            x1, y1, x2, y2,
            shape["left"], shape["top"],
            shape["right"], shape["bottom"],
        )
        if iou > best_iou and iou >= min_iou:
            best_iou = iou
            best_font = shape["font_name"]

    return best_font
